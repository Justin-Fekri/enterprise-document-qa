"""Generate sample enterprise PDF and Word documents for the demo corpus."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from enterprise_qa.config import SAMPLE_DOCS_DIR

NAVY = HexColor("#0F2744")
SLATE = HexColor("#334155")
RULE = HexColor("#94A3B8")
PALE = HexColor("#F1F5F9")


def generate_all(output_dir: Path | None = None) -> list[Path]:
    dest = output_dir or SAMPLE_DOCS_DIR
    dest.mkdir(parents=True, exist_ok=True)
    written = [
        _write_security_policy(dest / "information_security_policy.pdf"),
        _write_audit_procedure(dest / "internal_audit_procedure.pdf"),
        _write_risk_report(dest / "enterprise_risk_report.pdf"),
        _write_handbook(dest / "employee_handbook.docx"),
        _write_incident_plan(dest / "incident_response_plan.docx"),
    ]
    return written


def _styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CoverTitle",
            fontName="Times-Bold",
            fontSize=22,
            leading=26,
            textColor=NAVY,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CoverSub",
            fontName="Times-Italic",
            fontSize=11,
            leading=14,
            textColor=SLATE,
            spaceAfter=18,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H",
            fontName="Times-Bold",
            fontSize=13,
            leading=16,
            textColor=NAVY,
            spaceBefore=14,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            fontName="Times-Roman",
            fontSize=10.5,
            leading=14,
            textColor=SLATE,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BulletBody",
            fontName="Times-Roman",
            fontSize=10.5,
            leading=14,
            textColor=SLATE,
            leftIndent=16,
            spaceAfter=4,
        )
    )
    return styles


def _header_footer(canvas, doc, banner: str) -> None:
    canvas.saveState()
    canvas.setFillColor(NAVY)
    canvas.rect(0, LETTER[1] - 36, LETTER[0], 36, fill=1, stroke=0)
    canvas.setFillColor(white)
    canvas.setFont("Times-Bold", 9)
    canvas.drawString(0.75 * inch, LETTER[1] - 22, banner)
    canvas.setFillColor(RULE)
    canvas.rect(0, 0, LETTER[0], 32, fill=1, stroke=0)
    canvas.setFillColor(white)
    canvas.setFont("Times-Roman", 9)
    canvas.drawString(0.75 * inch, 14, "Northstar Holdings — Internal Use")
    canvas.drawRightString(LETTER[0] - 0.75 * inch, 14, f"Page {doc.page}")
    canvas.restoreState()


def _pdf(path: Path, banner: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.65 * inch,
    )


def _write_security_policy(path: Path) -> Path:
    styles = _styles()
    story = [
        Paragraph("Information Security Policy", styles["CoverTitle"]),
        Paragraph("Policy ID ISP-2026-01 · Effective 1 January 2026 · Classification: Confidential", styles["CoverSub"]),
        Paragraph("1. Purpose and scope", styles["H"]),
        Paragraph(
            "This policy establishes mandatory controls for protecting Northstar Holdings information "
            "assets, including customer records, financial data, source code, and audit workpapers. "
            "It applies to all employees, contractors, and vendors with network or document access. "
            "The Data Protection Officer is Morgan Ellis (dpo@northstar.example). Exceptions require "
            "written approval from the Chief Information Security Officer.",
            styles["Body"],
        ),
        Paragraph("2. Password and authentication standard", styles["H"]),
        Paragraph(
            "All interactive accounts must use a minimum password length of 14 characters, mixing "
            "upper and lower case letters, a numeral, and a symbol. Passwords must be rotated every "
            "90 days and must not reuse any of the prior 12 passwords. Shared service accounts are "
            "prohibited except where a vaulted credential is approved by Security Engineering. "
            "Multi-factor authentication is required for privileged access, remote VPN, and any "
            "system that stores Confidential or Restricted data. Passwords must never be stored in "
            "plain text, tickets, or chat messages.",
            styles["Body"],
        ),
        Paragraph("3. Data classification", styles["H"]),
        Paragraph(
            "Information is labeled Public, Internal, Confidential, or Restricted. Public information "
            "may be published externally. Internal information is for workforce use and includes the "
            "employee handbook and office procedures. Confidential information includes audit "
            "procedures, risk reports, and customer contracts. Restricted information includes "
            "incident response runbooks, encryption keys, and unpublished vulnerability findings. "
            "Restricted and Confidential data must not be sent through personal email, unsanctioned "
            "cloud drives, or removable media without encryption and CISO approval.",
            styles["Body"],
        ),
        PageBreak(),
        Paragraph("4. Access control", styles["H"]),
        Paragraph(
            "Access follows least privilege and role-based permissions. New access requests are "
            "ticketed, approved by the data owner, and provisioned within two business days. Access "
            "reviews occur quarterly for Confidential systems and monthly for Restricted systems. "
            "Terminated workforce members must have all access revoked within four hours of HR "
            "notification. Privileged sessions on production systems are recorded and retained for "
            "one year.",
            styles["Body"],
        ),
        Paragraph("5. Acceptable use", styles["H"]),
        Paragraph(
            "Company laptops must run the approved endpoint agent, full-disk encryption, and "
            "automatic screen lock after 10 minutes of inactivity. Software may be installed only "
            "from the internal catalog. Personal cloud backup of company documents is forbidden. "
            "Security event logging on servers is retained for 365 days. Suspected phishing must be "
            "reported to soc@northstar.example within one hour of discovery.",
            styles["Body"],
        ),
        Paragraph("6. Enforcement", styles["H"]),
        Paragraph(
            "Violations are investigated by Corporate Security and may result in disciplinary action "
            "up to termination and referral to law enforcement. This policy is reviewed annually or "
            "after a material incident. The current control owner is the CISO, reporting to the "
            "Audit and Risk Committee of the Board.",
            styles["Body"],
        ),
    ]
    doc = _pdf(path, "ISP-2026-01 · Information Security Policy")
    doc.build(story, onFirstPage=lambda c, d: _header_footer(c, d, "ISP-2026-01 · Information Security Policy"),
              onLaterPages=lambda c, d: _header_footer(c, d, "ISP-2026-01 · Information Security Policy"))
    return path


def _write_audit_procedure(path: Path) -> Path:
    styles = _styles()
    story = [
        Paragraph("Internal Audit Procedure", styles["CoverTitle"]),
        Paragraph("Procedure AUD-PROC-04 · Version 4.2 · Classification: Confidential", styles["CoverSub"]),
        Paragraph("1. Mandate", styles["H"]),
        Paragraph(
            "Internal Audit provides independent assurance to the Audit and Risk Committee that "
            "governance, risk management, and control processes are effective. The Chief Audit "
            "Executive is Priya Raman. Internal audits are conducted quarterly across in-scope "
            "processes, and an independent external financial statement audit is performed annually. "
            "Audit staff must remain independent of the activities they review and must disclose "
            "conflicts before fieldwork begins.",
            styles["Body"],
        ),
        Paragraph("2. Annual planning", styles["H"]),
        Paragraph(
            "Each November Internal Audit builds a risk-based plan using residual risk ratings, "
            "prior findings, regulatory change, and management requests. High residual-risk processes "
            "are audited at least annually. Medium-risk processes are audited at least every two "
            "years. The plan is approved by the Audit and Risk Committee before 15 December. "
            "Unplanned special reviews may be added with committee chair approval.",
            styles["Body"],
        ),
        Paragraph("3. Engagement lifecycle", styles["H"]),
        Paragraph(
            "A standard engagement has four phases: planning, fieldwork, reporting, and follow-up. "
            "Planning produces an engagement letter, process walkthrough, and test program. "
            "Fieldwork samples transactions, inspects evidence, and interviews process owners. "
            "Reporting issues a draft within ten business days of fieldwork close. Management has "
            "five business days to respond with action owners and dates. Final reports are issued "
            "to the committee and the responsible executive.",
            styles["Body"],
        ),
        PageBreak(),
        Paragraph("4. Finding severity and remediation SLAs", styles["H"]),
        Paragraph(
            "Findings are classified Critical, High, Medium, or Low. Critical findings indicate a "
            "control failure that could cause material financial misstatement, significant regulatory "
            "exposure, or a severe security incident. Critical findings must be remediated within "
            "15 calendar days. High findings must be remediated within 30 calendar days. Medium "
            "findings must be remediated within 90 calendar days. Low findings must be remediated "
            "within 180 calendar days. Overdue Critical or High items are escalated to the CEO and "
            "reported at the next committee meeting.",
            styles["Body"],
        ),
        Paragraph("5. Evidence standards", styles["H"]),
        Paragraph(
            "Working papers must identify the source, date obtained, and auditor. Screenshots require "
            "a timestamp and system name. Sampling uses a 95 percent confidence level for financial "
            "controls and a 90 percent confidence level for operational controls unless a statistical "
            "exception is documented. Evidence is retained for seven years in the audit workspace. "
            "Draft reports are Confidential; issued reports are Restricted until the committee "
            "releases a summary.",
            styles["Body"],
        ),
        Paragraph("6. Quality assurance", styles["H"]),
        Paragraph(
            "Every engagement is reviewed by a senior auditor who did not perform the tests. "
            "Internal Audit maintains an external quality assessment every five years in accordance "
            "with the Global Internal Audit Standards. Metrics reported quarterly include percent "
            "of plan complete, aging of open findings, and average days to issue a final report.",
            styles["Body"],
        ),
    ]
    banner = "AUD-PROC-04 · Internal Audit Procedure"
    doc = _pdf(path, banner)
    doc.build(story, onFirstPage=lambda c, d: _header_footer(c, d, banner),
              onLaterPages=lambda c, d: _header_footer(c, d, banner))
    return path


def _write_risk_report(path: Path) -> Path:
    styles = _styles()
    cell = ParagraphStyle("Cell", fontName="Times-Roman", fontSize=9, leading=12, textColor=SLATE)
    header_cell = ParagraphStyle("HCell", fontName="Times-Bold", fontSize=9, leading=12, textColor=white)
    rows = [
        [
            Paragraph("Risk", header_cell),
            Paragraph("Owner", header_cell),
            Paragraph("Inherent", header_cell),
            Paragraph("Residual", header_cell),
            Paragraph("Trend", header_cell),
        ],
        [
            Paragraph("Ransomware / destructive malware against core finance systems", cell),
            Paragraph("CISO", cell),
            Paragraph("Critical", cell),
            Paragraph("High", cell),
            Paragraph("Stable", cell),
        ],
        [
            Paragraph("Third-party SaaS outage affecting payroll and vendor payments", cell),
            Paragraph("CRO", cell),
            Paragraph("High", cell),
            Paragraph("Medium", cell),
            Paragraph("Improving", cell),
        ],
        [
            Paragraph("Regulatory reporting error under SOX and privacy statutes", cell),
            Paragraph("GC", cell),
            Paragraph("High", cell),
            Paragraph("Medium", cell),
            Paragraph("Stable", cell),
        ],
        [
            Paragraph("Key-person concentration in treasury operations", cell),
            Paragraph("CFO", cell),
            Paragraph("Medium", cell),
            Paragraph("Medium", cell),
            Paragraph("Worsening", cell),
        ],
    ]
    table = Table(rows, colWidths=[2.8 * inch, 0.8 * inch, 0.9 * inch, 0.9 * inch, 0.9 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("BACKGROUND", (0, 1), (-1, -1), PALE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.3, RULE),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story = [
        Paragraph("Enterprise Risk Report — Q1 2026", styles["CoverTitle"]),
        Paragraph("Prepared by Enterprise Risk Management · Classification: Confidential", styles["CoverSub"]),
        Paragraph("1. Executive summary", styles["H"]),
        Paragraph(
            "Northstar Holdings operates a 5×5 likelihood-impact risk matrix. Residual risk is "
            "accepted only at Medium or below unless the Board grants a time-bound exception. "
            "The Chief Risk Officer is Dana Voss. This quarter the top residual exposure remains "
            "ransomware against core finance systems, rated High after compensating controls. "
            "Third-party vendor risk is Medium following dual-provider payroll cutover. No new "
            "Critical residual risks were accepted in Q1 2026.",
            styles["Body"],
        ),
        Paragraph("2. Top residual risks", styles["H"]),
        table,
        Spacer(1, 12),
        Paragraph("3. Ransomware scenario", styles["H"]),
        Paragraph(
            "The ransomware residual risk rating is High. Inherent risk is Critical because a "
            "successful encryption event could halt close, payroll, and customer billing for more "
            "than 72 hours. Compensating controls include immutable backups tested monthly, "
            "network segmentation of finance servers, privileged access workstations, and a "
            "tabletop exercise completed in February 2026. Residual risk remains High because "
            "backup restore for the general ledger still exceeds the 8-hour recovery time objective. "
            "The treatment plan is to complete ledger restore automation by 30 June 2026, owned "
            "by the CISO with Finance Operations as the business sponsor.",
            styles["Body"],
        ),
        PageBreak(),
        Paragraph("4. Third-party and vendor risk", styles["H"]),
        Paragraph(
            "Third-party vendor residual risk is Medium. All vendors that process Confidential data "
            "require SOC 2 Type II or equivalent, annual security questionnaires, and a right-to-audit "
            "clause. Critical vendors are reviewed twice per year. A payroll provider outage in "
            "November 2025 delayed one cycle by 19 hours; a secondary provider is now contracted. "
            "Vendors without current SOC 2 reports are not permitted to receive Restricted data.",
            styles["Body"],
        ),
        Paragraph("5. Risk appetite and reporting", styles["H"]),
        Paragraph(
            "Risk appetite statements: (a) zero tolerance for intentional regulatory breaches; "
            "(b) no more than two High residual technology risks without a Board-dated treatment "
            "plan; (c) recovery time for material financial systems must not exceed 8 hours after "
            "the treatment plan closes. KRIs reported monthly to the committee include phishing "
            "fail rate, privileged account count, open Critical audit findings, and backup restore "
            "success. This report is Confidential and is not for distribution outside the committee, "
            "executive team, Internal Audit, and Compliance.",
            styles["Body"],
        ),
    ]
    banner = "ERR-2026-Q1 · Enterprise Risk Report"
    doc = _pdf(path, banner)
    doc.build(story, onFirstPage=lambda c, d: _header_footer(c, d, banner),
              onLaterPages=lambda c, d: _header_footer(c, d, banner))
    return path


def _docx_heading(document: Document, text: str) -> None:
    heading = document.add_heading(text, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(15, 39, 68)


def _write_handbook(path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    title = document.add_heading("Northstar Holdings Employee Handbook", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(15, 39, 68)
    lead = document.add_paragraph()
    run = lead.add_run("Effective 1 January 2026  ·  Classification: Internal  ·  Applies to all employees")
    run.italic = True
    run.font.size = Pt(11)

    _docx_heading(document, "Conduct and workplace")
    document.add_paragraph(
        "Northstar Holdings expects professional conduct, accurate time reporting, and respectful "
        "collaboration. Harassment, discrimination, and retaliation are prohibited. Concerns may be "
        "raised to a manager, Human Resources, or the anonymous ethics line at ethics@northstar.example. "
        "Retaliation against good-faith reporters is itself a policy violation."
    )
    _docx_heading(document, "Working hours, leave, and hybrid work")
    document.add_paragraph(
        "Standard hours are 40 per week. Employees receive 20 days of paid time off, 10 company "
        "holidays, and up to 12 weeks of parental leave. Hybrid employees may work remotely up to "
        "three days per week with manager approval. Core collaboration hours are 10:00–15:00 local "
        "time. Overtime for non-exempt staff must be pre-approved. Company devices must be used for "
        "work email; personal email is not an approved records system."
    )
    _docx_heading(document, "Handling company information")
    document.add_paragraph(
        "Employees may access Internal documents such as this handbook, office maps, and published "
        "benefits guides. Confidential risk reports, audit procedures, and Restricted incident "
        "runbooks are limited to named roles. Restricted and Confidential data must not be sent "
        "through personal email. Screenshots of customer data are prohibited. If you are unsure of "
        "a classification, treat the material as Confidential and ask your manager or Compliance."
    )
    _docx_heading(document, "Security expectations for every employee")
    document.add_paragraph(
        "Lock your workstation when you step away. Do not share passwords. Complete quarterly "
        "security awareness training by the assigned deadline. Report suspected phishing to "
        "soc@northstar.example. The minimum password length is 14 characters, and passwords must "
        "be rotated every 90 days. Lost badges or laptops must be reported within two hours. "
        "Visitors must be signed in at reception and escorted in Restricted areas."
    )
    _docx_heading(document, "Performance and records")
    document.add_paragraph(
        "Performance reviews occur in June and December. Personnel files are Internal and visible "
        "to the employee, their manager, and HR. Compensation data is Confidential. This handbook "
        "does not create a contract of employment. Northstar Holdings is an at-will employer except "
        "where a written agreement states otherwise. Questions: hr@northstar.example."
    )
    document.save(str(path))
    return path


def _write_incident_plan(path: Path) -> Path:
    document = Document()
    title = document.add_heading("Cybersecurity Incident Response Plan", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(15, 39, 68)
    lead = document.add_paragraph()
    run = lead.add_run("Plan IRP-2026-A  ·  Classification: Restricted  ·  Owner: CISO")
    run.italic = True

    _docx_heading(document, "Purpose")
    document.add_paragraph(
        "This plan describes how Northstar Holdings detects, contains, eradicates, and recovers "
        "from cybersecurity incidents. It is Restricted because it documents detection sources, "
        "escalation paths, and decision criteria. Distribution is limited to Security Operations, "
        "the CISO office, Compliance, and named executives. Do not forward this document."
    )
    _docx_heading(document, "Severity model")
    document.add_paragraph(
        "Incidents are graded P1 through P4. A P1 incident is an active compromise of production "
        "finance, identity, or customer systems, ransomware encryption, or confirmed exfiltration "
        "of Restricted data. P1 incidents require a 15-minute response from Security Operations "
        "and a 60-minute executive briefing. P2 incidents are confirmed intrusions with limited "
        "blast radius and require a 1-hour response. P3 incidents are suspicious activity needing "
        "investigation within 4 hours. P4 incidents are policy or malware events on a single "
        "endpoint with no evidence of lateral movement and require a next-business-day response."
    )
    _docx_heading(document, "Operating model")
    document.add_paragraph(
        "The Security Operations Center is staffed 24 hours a day, 7 days a week. The incident "
        "commander for P1 and P2 events is the on-call SOC lead, who may escalate to the CISO. "
        "Legal, Communications, and the Data Protection Officer (Morgan Ellis) join the bridge for "
        "any incident involving personal data. Containment actions that take a material system "
        "offline require CFO concurrence except where delay would worsen a P1 event."
    )
    _docx_heading(document, "Notification and evidence")
    document.add_paragraph(
        "Customer or regulator notification decisions are made by Legal within 24 hours of P1 "
        "declaration when personal data may be affected. Forensic images are captured before "
        "reimaging. Chat and bridge recordings for P1 events are retained for two years. After-action "
        "reviews are completed within 10 business days. Tabletop exercises run twice per year; the "
        "last exercise was completed in February 2026 against a ransomware scenario."
    )
    _docx_heading(document, "Return to operations")
    document.add_paragraph(
        "Recovery is complete when the recovered system meets the 8-hour recovery time objective "
        "for material financial systems, malware is eradicated, credentials are rotated, and "
        "monitoring rules for the observed tactics are deployed. The CISO closes the incident in "
        "the tracker. Lessons learned that require control changes are entered as audit-style "
        "actions with owners and dates."
    )
    document.save(str(path))
    return path


if __name__ == "__main__":
    paths = generate_all()
    for path in paths:
        print(path)
