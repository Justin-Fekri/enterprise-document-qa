"""Round-trip real PDF and DOCX files through the loaders."""

from pathlib import Path

import pytest

from app.ingestion.loaders import (
    UnsupportedDocument,
    _looks_like_heading,
    load_document,
    normalize,
)

pytest.importorskip("reportlab", reason="reportlab is needed to build test PDFs")


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    path = tmp_path / "sample.pdf"
    styles = getSampleStyleSheet()
    SimpleDocTemplate(str(path), pagesize=LETTER).build(
        [
            Paragraph("4.2 Access Control", styles["Heading2"]),
            Paragraph("Privileged sessions expire after four hours.", styles["BodyText"]),
            PageBreak(),
            Paragraph("5.1 Incident Response", styles["Heading2"]),
            Paragraph("Incidents must be reported within one hour.", styles["BodyText"]),
        ]
    )
    return path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    import docx
    from docx.enum.text import WD_BREAK

    path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_heading("Audit Procedure", level=1)
    document.add_paragraph("Sample sizes follow the attribute sampling table.")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Findings", level=1)
    document.add_paragraph("Critical findings escalate within five business days.")

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Risk ID"
    table.rows[0].cells[1].text = "R-2024-017"
    document.save(str(path))
    return path


def test_pdf_pages_are_numbered_from_one(sample_pdf):
    blocks = load_document(sample_pdf)
    assert [block.page for block in blocks] == [1, 2]
    assert "four hours" in blocks[0].text
    assert "within one hour" in blocks[1].text


def test_pdf_headings_are_captured(sample_pdf):
    blocks = load_document(sample_pdf)
    assert blocks[0].section.startswith("4.2")


def test_docx_pages_split_on_explicit_page_breaks(sample_docx):
    blocks = load_document(sample_docx)
    pages = {block.page for block in blocks}
    assert len(pages) >= 2
    assert any("attribute sampling" in block.text for block in blocks)
    assert any("five business days" in block.text for block in blocks)


def test_docx_tables_are_extracted(sample_docx):
    blocks = load_document(sample_docx)
    assert any("R-2024-017" in block.text for block in blocks)


def test_unsupported_type_is_rejected(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello")
    with pytest.raises(UnsupportedDocument):
        load_document(path)


def test_normalize_repairs_hyphenated_line_wraps():
    assert "authentication" in normalize("authen-\ntication required")


def test_heading_detection():
    assert _looks_like_heading("4.2 Access Control")
    assert _looks_like_heading("INCIDENT RESPONSE")
    assert not _looks_like_heading("This is an ordinary sentence of body text.")
