"""Load PDF and Word documents into page-anchored text blocks.

Both loaders return the same shape - `(page_number, section_heading, text)` -
so the chunker never needs to know which format a document came from. Page
numbers are what make page-level citations possible downstream.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx"}

# Word stores an explicit page break as a <w:br w:type="page"/> run.
_PAGE_BREAK_XML = 'w:type="page"'


@dataclass
class PageBlock:
    """One page (PDF) or one page-break-delimited block (DOCX)."""

    page: int
    section: str
    text: str


class UnsupportedDocument(ValueError):
    """Raised for file types outside the PDF/Word scope of this assistant."""


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(65536), b""):
            digest.update(block)
    return digest.hexdigest()[:16]


def normalize(text: str) -> str:
    """Collapse the whitespace noise typical of PDF text extraction."""
    text = text.replace("­", "")  # soft hyphens
    text = re.sub(r"-\n(?=[a-z])", "", text)  # de-hyphenate line wraps
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not (3 < len(line) < 90):
        return False
    if line.endswith("."):
        return False
    if re.match(r"^\d+(\.\d+)*\s+\S", line):  # "4.2 Access Control"
        return True
    words = line.split()
    return len(words) <= 10 and (line.isupper() or line.istitle())


def load_pdf(path: Path) -> list[PageBlock]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[PageBlock] = []
    current_section = ""

    for index, page in enumerate(reader.pages, start=1):
        raw = normalize(page.extract_text() or "")
        if not raw:
            continue
        for line in raw.split("\n")[:5]:
            if _looks_like_heading(line):
                current_section = line.strip()
                break
        blocks.append(PageBlock(page=index, section=current_section, text=raw))

    return blocks


def load_docx(path: Path) -> list[PageBlock]:
    """Word has no fixed pagination, so we page on explicit page breaks.

    Documents without any page breaks fall back to one block per ~40
    paragraphs, which keeps citation targets small enough to be useful.
    """
    import docx

    document = docx.Document(str(path))
    blocks: list[PageBlock] = []
    page = 1
    section = ""
    buffer: list[str] = []
    paragraphs_since_break = 0

    def flush() -> None:
        nonlocal buffer, paragraphs_since_break
        text = normalize("\n".join(buffer))
        if text:
            blocks.append(PageBlock(page=page, section=section, text=text))
        buffer = []
        paragraphs_since_break = 0

    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "") if paragraph.style else ""
        if style.startswith("Heading") and paragraph.text.strip():
            section = paragraph.text.strip()

        if paragraph.text.strip():
            buffer.append(paragraph.text.strip())
            paragraphs_since_break += 1

        has_break = _PAGE_BREAK_XML in paragraph._p.xml
        if has_break or paragraphs_since_break >= 40:
            flush()
            page += 1

    flush()

    # Tables carry a lot of the substance in audit and risk documents.
    for table in document.tables:
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]
        if rows:
            blocks.append(
                PageBlock(page=page, section=section or "Table", text="\n".join(rows))
            )

    return blocks


def load_document(path: Path) -> list[PageBlock]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    raise UnsupportedDocument(
        f"{path.name}: only {', '.join(sorted(SUPPORTED_SUFFIXES))} are supported"
    )
