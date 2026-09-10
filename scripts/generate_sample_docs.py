"""Generate a small synthetic corpus of business documents.

Everything here describes a fictional company, "Northwind Systems". The
content is invented for demonstration and evaluation only - it is not advice
and does not describe any real organisation's controls.

Run:  python scripts/generate_sample_docs.py
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.config import get_settings  # noqa: E402

SECURITY_POLICY = [
    (
        "1. Purpose and Scope",
        [
            "This Information Security Policy defines the minimum security "
            "requirements for all Northwind Systems information assets, "
            "employees, contractors and third-party service providers.",
            "The policy applies to all systems that store, process or transmit "
            "Northwind data, including cloud services and employee-owned "
            "devices enrolled in the mobile device management programme.",
            "Exceptions to this policy must be approved in writing by the Chief "
            "Information Security Officer and are valid for no more than 90 days.",
        ],
    ),
    (
        "2. Password and Authentication Requirements",
        [
            "All user account passwords must be a minimum of 14 characters and "
            "must not appear in the maintained list of breached credentials.",
            "Standard user passwords must be rotated every 180 days. Privileged "
            "and administrative account passwords must be rotated every 90 days.",
            "Multi-factor authentication is mandatory for all remote access, all "
            "administrative consoles and all access to systems classified as "
            "Confidential or Restricted.",
            "Accounts are locked after 5 consecutive failed authentication "
            "attempts and remain locked for 30 minutes or until reset by the "
            "service desk.",
            "Shared or generic accounts are prohibited except for documented "
            "service accounts, which must use credentials of at least 32 "
            "characters stored in the enterprise secrets vault.",
        ],
    ),
    (
        "3. Data Classification",
        [
            "Northwind uses four classification levels: Public, Internal, "
            "Confidential and Restricted.",
            "Restricted data includes customer payment information, employee "
            "medical records, and unpublished financial results. Restricted data "
            "must be encrypted at rest using AES-256 and in transit using TLS 1.2 "
            "or higher.",
            "Confidential data must not be copied to removable media without a "
            "documented business justification approved by a data owner.",
            "Data retention for Internal records is 3 years. Restricted records "
            "are retained for 7 years and then securely destroyed.",
        ],
    ),
    (
        "4. Incident Response",
        [
            "All suspected security incidents must be reported to the Security "
            "Operations Centre within 1 hour of discovery.",
            "The incident response team must triage and assign a severity rating "
            "within 4 hours of a report being received.",
            "Severity 1 incidents, defined as confirmed unauthorised access to "
            "Restricted data, require notification of the executive team within "
            "12 hours and a written post-incident review within 10 business days.",
            "Regulatory notification, where required, is coordinated by the Legal "
            "team and must occur within 72 hours of incident confirmation.",
        ],
    ),
    (
        "5. Vendor and Third-Party Security",
        [
            "Third parties handling Confidential or Restricted data must complete "
            "a security assessment before contract signature and annually "
            "thereafter.",
            "Vendors must notify Northwind of any security breach affecting "
            "Northwind data within 24 hours of detection.",
            "Vendor access is provisioned as time-bound accounts that expire "
            "automatically after 12 months unless renewed.",
        ],
    ),
    (
        "6. Security Awareness Training",
        [
            "All employees must complete security awareness training within 30 "
            "days of their start date and annually thereafter.",
            "Employees in engineering roles must additionally complete secure "
            "development training covering the OWASP Top Ten every 12 months.",
            "Simulated phishing exercises are conducted quarterly. Employees who "
            "fail two consecutive exercises are enrolled in remedial training.",
        ],
    ),
]

ACCESS_CONTROL_STANDARD = [
    (
        "1. Overview",
        [
            "This Access Control Standard implements Section 2 of the Information "
            "Security Policy and is classified Restricted because it documents "
            "control weaknesses and compensating measures.",
            "The standard covers identity lifecycle, privileged access and "
            "periodic access review for all production systems.",
        ],
    ),
    (
        "2. Identity Lifecycle",
        [
            "User accounts are provisioned from the HR system of record. Access "
            "requests outside the standard role profile require line manager and "
            "system owner approval.",
            "Accounts for departing employees must be disabled within 4 hours of "
            "the termination record being created in the HR system.",
            "Dormant accounts with no authentication activity for 60 days are "
            "automatically disabled.",
        ],
    ),
    (
        "3. Privileged Access Management",
        [
            "Privileged access is granted just-in-time through the privileged "
            "access management platform, with a maximum session duration of 4 "
            "hours.",
            "All privileged sessions are recorded and the recordings retained for "
            "12 months.",
            "Break-glass accounts exist for two production environments. Their "
            "credentials are held in sealed escrow and any use triggers an "
            "immediate page to the on-call security engineer.",
            "Known gap: three legacy manufacturing systems do not support the "
            "privileged access platform. Compensating control is a manually "
            "reviewed monthly access log, tracked as risk item R-2024-017.",
        ],
    ),
    (
        "4. Access Review",
        [
            "System owners must certify all user entitlements for in-scope "
            "systems every quarter.",
            "Entitlements not certified within 15 business days of the review "
            "opening are automatically revoked.",
            "Segregation of duties conflicts identified during review must be "
            "remediated or formally accepted within 30 days.",
        ],
    ),
]


def _write_pdf(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    heading = ParagraphStyle(
        "SectionHeading", parent=styles["Heading2"], spaceBefore=14, spaceAfter=8
    )
    body = ParagraphStyle(
        "BodyTextJustified", parent=styles["BodyText"], leading=15, spaceAfter=8
    )

    story = [
        Paragraph(title, styles["Title"]),
        Paragraph(
            "Northwind Systems Inc. - synthetic sample document for demonstration.",
            styles["Italic"],
        ),
        Spacer(1, 0.3 * inch),
    ]

    for index, (name, paragraphs) in enumerate(sections):
        # One section per page keeps citation page numbers meaningful.
        if index:
            story.append(PageBreak())
        story.append(Paragraph(name, heading))
        story.extend(Paragraph(text, body) for text in paragraphs)

    SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        title=title,
        author="Northwind Systems",
        topMargin=0.9 * inch,
        bottomMargin=0.9 * inch,
    ).build(story)


def _write_docx(
    path: Path,
    title: str,
    sections: list[tuple[str, list[str]]],
    table: tuple[list[str], list[list[str]]] | None = None,
) -> None:
    import docx
    from docx.enum.text import WD_BREAK

    document = docx.Document()
    document.add_heading(title, level=0)
    document.add_paragraph(
        "Northwind Systems Inc. - synthetic sample document for demonstration."
    )

    for index, (name, paragraphs) in enumerate(sections):
        if index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading(name, level=1)
        for text in paragraphs:
            document.add_paragraph(text)

    if table is not None:
        header, rows = table
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_heading("Risk Register", level=1)
        grid = document.add_table(rows=1, cols=len(header))
        grid.style = "Table Grid"
        for cell, value in zip(grid.rows[0].cells, header, strict=False):
            cell.text = value
        for row in rows:
            cells = grid.add_row().cells
            for cell, value in zip(cells, row, strict=False):
                cell.text = value

    document.save(str(path))


AUDIT_PROCEDURE = [
    (
        "1. Audit Charter",
        [
            "The Internal Audit function provides independent assurance over the "
            "design and operating effectiveness of Northwind's internal controls.",
            "Internal Audit reports functionally to the Audit Committee and "
            "administratively to the Chief Financial Officer.",
            "The annual audit plan is approved by the Audit Committee before the "
            "start of each fiscal year and is refreshed at the half year.",
        ],
    ),
    (
        "2. Engagement Procedure",
        [
            "Each engagement follows five phases: planning, walkthrough, testing, "
            "reporting and follow-up.",
            "A planning memorandum documenting scope, objectives and the control "
            "population must be issued to the auditee at least 10 business days "
            "before fieldwork begins.",
            "Sample sizes follow the standard attribute sampling table: 25 items "
            "for controls operating daily, 15 for weekly controls, 5 for monthly "
            "controls and 2 for quarterly controls.",
            "Evidence must be retained in the audit management system for 7 years "
            "and must be sufficient for an independent reviewer to re-perform the "
            "test.",
        ],
    ),
    (
        "3. Findings and Ratings",
        [
            "Findings are rated Critical, High, Medium or Low based on financial, "
            "regulatory and reputational impact.",
            "Critical findings must be reported to the Audit Committee chair "
            "within 5 business days of being confirmed with management.",
            "Management must provide a remediation plan with a named owner and a "
            "target date within 20 business days of receiving the draft report.",
            "The final report is issued within 30 business days of the close of "
            "fieldwork.",
        ],
    ),
    (
        "4. Follow-Up and Closure",
        [
            "Internal Audit validates remediation evidence before a finding is "
            "closed. Self-attestation by management is not sufficient for "
            "Critical or High findings.",
            "Overdue findings are escalated to the Audit Committee at each "
            "quarterly meeting.",
            "A finding may be closed as risk-accepted only with written approval "
            "from the accountable executive and the Chief Risk Officer.",
        ],
    ),
]

RISK_REPORT = [
    (
        "1. Executive Summary",
        [
            "This report summarises Northwind's enterprise risk position for the "
            "third quarter. The aggregate residual risk score decreased from 62 "
            "to 57 quarter over quarter.",
            "Two risks moved above appetite: third-party concentration risk and "
            "legacy system obsolescence. Both have funded remediation plans.",
            "No new Critical risks were identified during the quarter.",
        ],
    ),
    (
        "2. Risk Appetite",
        [
            "The Board has set an aggregate residual risk appetite threshold of "
            "60 on the 100-point enterprise scale.",
            "Individual risks scoring above 15 residual are reported to the Board "
            "Risk Committee each quarter regardless of trend.",
            "Operational risk appetite permits a maximum of 4 hours of unplanned "
            "downtime per quarter for tier-one customer-facing services.",
        ],
    ),
    (
        "3. Key Risk Movements",
        [
            "Third-party concentration risk increased from 14 to 18 residual "
            "following the consolidation of payment processing with a single "
            "provider. A secondary provider is being onboarded, with completion "
            "targeted for the end of the fourth quarter.",
            "Legacy system obsolescence increased from 15 to 17 residual. Three "
            "manufacturing systems remain outside the privileged access platform, "
            "tracked as R-2024-017.",
            "Cyber intrusion risk decreased from 21 to 16 residual after "
            "multi-factor authentication was extended to all remote access.",
        ],
    ),
]

RISK_TABLE = (
    ["ID", "Risk", "Inherent", "Residual", "Owner", "Status"],
    [
        ["R-2024-011", "Third-party concentration", "24", "18", "COO", "Above appetite"],
        ["R-2024-017", "Legacy system obsolescence", "22", "17", "CIO", "Above appetite"],
        ["R-2024-004", "Cyber intrusion", "27", "16", "CISO", "Within appetite"],
        ["R-2024-022", "Regulatory reporting error", "18", "11", "CFO", "Within appetite"],
        ["R-2024-030", "Key person dependency", "14", "9", "CHRO", "Within appetite"],
    ],
)


def main() -> None:
    settings = get_settings()
    target = settings.document_dir
    target.mkdir(parents=True, exist_ok=True)

    _write_pdf(
        target / "information_security_policy.pdf",
        "Information Security Policy",
        SECURITY_POLICY,
    )
    _write_pdf(
        target / "access_control_standard.pdf",
        "Access Control Standard",
        ACCESS_CONTROL_STANDARD,
    )
    _write_docx(
        target / "internal_audit_procedure.docx",
        "Internal Audit Procedure Manual",
        AUDIT_PROCEDURE,
    )
    _write_docx(
        target / "q3_risk_report.docx",
        "Q3 Enterprise Risk Report",
        RISK_REPORT,
        table=RISK_TABLE,
    )

    for path in sorted(target.iterdir()):
        print(f"  wrote {path.relative_to(settings.document_dir.parent.parent)}")


if __name__ == "__main__":
    main()
